"""Editing an existing Word document safely (backend/utils/docx_edit.py): a verified edited COPY, never the original."""
import copy
import io
import zipfile

import pytest
from docx import Document
from lxml import etree

from backend.utils import docx_edit as de

W = de.W


def _q(tag):
    return f"{{{W}}}{tag}"


def _fixture(extra=None):
    """A realistic document: a title, mixed-format runs, a table, a header, a list, plus whatever `extra(document)` adds."""
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Confidential header"
    doc.add_heading("Jack Harper", level=0)                                   # 1
    para = doc.add_paragraph()                                                # 2
    para.add_run("Software engineer at ")
    para.add_run("Principal").bold = True
    para.add_run(" since 2023.")
    doc.add_paragraph("Skills: Python, FastAPI, React.")                      # 3
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Role", "Years"           # 4, 5
    table.cell(1, 0).text, table.cell(1, 1).text = "Engineer", "3"           # 6, 7
    doc.add_paragraph("Built Sonic.", style="List Bullet")                    # 8
    doc.add_paragraph("Last line.")                                           # 9
    if extra:
        extra(doc)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _members(data):
    archive = zipfile.ZipFile(io.BytesIO(data))
    return {info.filename: archive.read(info.filename) for info in archive.infolist()}


def _texts(data):
    return [p.text for p in Document(io.BytesIO(data)).paragraphs]


def _outline(data):
    return de.parse_outline(de.outline_text(data))


def _inject_paragraph(doc, xml):
    body = doc.element.body
    element = etree.fromstring(f'<w:p xmlns:w="{W}">{xml}</w:p>')
    body.insert(len(body) - 1, element)         # before the final section properties


# ---- the outline Sonic reads ----------------------------------------------------------------------------------------------------------------------

def test_the_outline_numbers_every_paragraph_in_order_with_table_context_and_marks_the_ones_that_cannot_be_deleted():
    text = de.outline_text(_fixture())
    assert text.startswith("WORD DOCUMENT") and "[1] Jack Harper" in text
    assert "[2] Software engineer at Principal since 2023." in text
    assert "[4~|table 1 row 1 col 1] Role" in text and "[7~|table 1 row 2 col 2] 3" in text      # the flag comes straight after the number, then the table position
    outline = de.parse_outline(text)
    assert outline[2]["text"] == "Software engineer at Principal since 2023." and outline[2]["deletable"]
    assert outline[4] == {"editable": True, "deletable": False, "insertable": True, "text": "Role"}      # the only paragraph in its cell
    assert [outline[n]["text"] for n in (6, 7)] == ["Engineer", "3"]


def test_text_with_line_breaks_and_backslashes_round_trips_through_the_outline():
    def extra(doc):
        paragraph = doc.add_paragraph("first")
        paragraph.add_run().add_break()
        paragraph.add_run("second \\ end")
    outline = _outline(_fixture(extra))
    last = outline[max(outline)]
    assert last["text"] == "first\nsecond \\ end"


@pytest.mark.parametrize("data, fragment", [
    (b"", "isn't a Word document"),
    (b"not a zip at all", "valid .docx"),
    (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 100, "password-protected or an old .doc"),
    (b"PK" + b"0" * 30, "valid .docx"),
])
def test_files_that_are_not_documents_are_refused_with_a_reason(data, fragment):
    with pytest.raises(de.DocxEditError, match=fragment):
        de.outline_text(data)


def test_a_document_with_macros_or_a_doctype_or_no_body_or_too_big_is_refused():
    def rebuild(change):
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as archive:
            for name, content in change(_members(_fixture())).items():
                archive.writestr(name, content)
        return out.getvalue()

    with pytest.raises(de.DocxEditError, match="macros"):
        de.outline_text(rebuild(lambda m: {**m, "word/vbaProject.bin": b"x"}))
    with pytest.raises(de.DocxEditError, match="can't be edited safely"):
        de.outline_text(rebuild(lambda m: {**m, "word/document.xml": b'<!DOCTYPE x [<!ENTITY a "b">]>' + m["word/document.xml"]}))
    with pytest.raises(de.DocxEditError, match="no document body"):
        de.outline_text(rebuild(lambda m: {k: v for k, v in m.items() if k != "word/document.xml"}))
    with pytest.raises(de.DocxEditError, match="too large"):
        de.outline_text(b"PK" + b"0" * (de.MAX_DOCX_BYTES + 1))


# ---- paragraphs that must not be edited ----------------------------------------------------------------------------------------------------------

def test_tracked_changes_links_fields_and_content_controls_are_read_but_never_editable():
    def extra(doc):
        _inject_paragraph(doc, '<w:ins w:id="1" w:author="x" w:date="2026-01-01T00:00:00Z"><w:r><w:t>tracked</w:t></w:r></w:ins>')
        _inject_paragraph(doc, '<w:hyperlink w:anchor="a"><w:r><w:t>a link</w:t></w:r></w:hyperlink>')
        _inject_paragraph(doc, '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText>PAGE</w:instrText></w:r><w:r><w:t>1</w:t></w:r>')
        _inject_paragraph(doc, '<w:r><w:rPr><w:rPrChange w:id="2" w:author="x"><w:rPr/></w:rPrChange></w:rPr><w:t>fmt change</w:t></w:r>')
        doc.element.body.insert(len(doc.element.body) - 1, etree.fromstring(
            f'<w:sdt xmlns:w="{W}"><w:sdtContent><w:p><w:r><w:t>in a control</w:t></w:r></w:p></w:sdtContent></w:sdt>'))
    data = _fixture(extra)
    outline = _outline(data)
    flagged = {n: v for n, v in outline.items() if not v["editable"]}
    assert [v["text"] for v in flagged.values()] == ["tracked", "a link", "1", "fmt change", "in a control"]
    assert outline[max(outline)]["insertable"] is False                      # the one inside the content control
    n = next(n for n, v in flagged.items() if v["text"] == "a link")
    with pytest.raises(de.DocxEditError, match="can't be edited safely"):
        de.apply_edits(data, [{"op": "replace", "paragraph": n, "old": "a link", "new": "x"}])
    with pytest.raises(de.DocxEditError, match="can't be deleted"):
        de.apply_edits(data, [{"op": "delete", "paragraph": n}])
    last = max(outline)
    with pytest.raises(de.DocxEditError, match="can't be inserted"):
        de.apply_edits(data, [{"op": "insert_after", "paragraph": last, "text": "x"}])


# ---- the edit itself --------------------------------------------------------------------------------------------------------------------------------

def test_replacing_text_inside_one_run_changes_only_those_words_and_keeps_every_other_run_and_its_formatting():
    data = _fixture()
    edited = de.apply_edits(data, [{"op": "replace", "paragraph": 3, "old": "FastAPI", "new": "FastAPI, Postgres"}])
    texts = _texts(edited)
    assert "Skills: Python, FastAPI, Postgres, React." in texts
    assert [t for t in _texts(data) if t != "Skills: Python, FastAPI, React."] == [t for t in texts if t != "Skills: Python, FastAPI, Postgres, React."]


def test_replacing_text_across_formatted_runs_keeps_formatting_either_side_and_gives_new_text_the_first_characters_style():
    data = _fixture()
    edited = de.apply_edits(data, [{"op": "replace", "paragraph": 2, "old": "at Principal since", "new": "at Acme from"}])
    paragraph = Document(io.BytesIO(edited)).paragraphs[1]
    assert paragraph.text == "Software engineer at Acme from 2023."
    bold = [run.text for run in paragraph.runs if run.bold]
    assert bold == []                       # "at " (plain) started the span, so the new words are plain; "Principal" (bold) was inside the replaced span
    assert paragraph.runs[0].text.startswith("Software engineer ")
    assert paragraph.runs[-1].text.endswith(" 2023.")


def test_a_replacement_that_deletes_a_whole_bold_run_removes_it_cleanly():
    edited = de.apply_edits(_fixture(), [{"op": "replace", "paragraph": 2, "old": "Principal", "new": ""}])
    paragraph = Document(io.BytesIO(edited)).paragraphs[1]
    assert paragraph.text == "Software engineer at  since 2023." and not any(run.bold for run in paragraph.runs)


def test_several_edits_all_refer_to_the_numbers_in_the_outline_that_was_read_so_they_do_not_shift_each_other():
    data = _fixture()
    edited = de.apply_edits(data, [
        {"op": "delete", "paragraph": 3},
        {"op": "insert_after", "paragraph": 2, "text": "New summary line."},
        {"op": "replace", "paragraph": 9, "old": "Last", "new": "Final"},
        {"op": "insert_after", "paragraph": 2, "text": "Second new line."},
    ])
    texts = _texts(edited)
    assert texts[:4] == ["Jack Harper", "Software engineer at Principal since 2023.", "New summary line.", "Second new line."]
    assert "Skills: Python, FastAPI, React." not in texts and texts[-1] == "Final line."
    assert len(texts) == len(_texts(data)) + 2 - 1


def test_an_inserted_paragraph_copies_the_paragraph_formatting_of_the_one_it_follows():
    edited = de.apply_edits(_fixture(), [{"op": "insert_after", "paragraph": 8, "text": "Another bullet."}])
    paragraphs = Document(io.BytesIO(edited)).paragraphs
    index = [p.text for p in paragraphs].index("Another bullet.")
    assert paragraphs[index].style.name == "List Bullet" == paragraphs[index - 1].style.name


def test_inserting_into_a_table_cell_adds_a_paragraph_in_that_cell():
    edited = de.apply_edits(_fixture(), [{"op": "insert_after", "paragraph": 6, "text": "Also this"}])
    cell = Document(io.BytesIO(edited)).tables[0].cell(1, 0)
    assert [p.text for p in cell.paragraphs] == ["Engineer", "Also this"]


def test_the_only_paragraph_in_a_table_cell_cannot_be_deleted_but_its_text_can_change():
    with pytest.raises(de.DocxEditError, match="only paragraph in its table cell"):
        de.apply_edits(_fixture(), [{"op": "delete", "paragraph": 6}])
    edited = de.apply_edits(_fixture(), [{"op": "replace", "paragraph": 7, "old": "3", "new": "4"}])
    assert Document(io.BytesIO(edited)).tables[0].cell(1, 1).text == "4"


def test_bookmarks_and_spellcheck_marks_stay_where_they_were_in_an_edited_paragraph():
    def extra(doc):
        _inject_paragraph(doc, '<w:bookmarkStart w:id="9" w:name="mark"/><w:r><w:t>Before mark</w:t></w:r><w:proofErr w:type="spellStart"/>'
                               '<w:r><w:t xml:space="preserve"> after mark</w:t></w:r><w:bookmarkEnd w:id="9"/>')
    data = _fixture(extra)
    number = max(_outline(data))
    edited = de.apply_edits(data, [{"op": "replace", "paragraph": number, "old": "after", "new": "past"}])
    root = etree.fromstring(_members(edited)["word/document.xml"])
    last = list(root.iter(_q("p")))[number - 1]
    assert [etree.QName(child).localname for child in last] == ["bookmarkStart", "r", "proofErr", "r", "bookmarkEnd"]
    assert Document(io.BytesIO(edited)).paragraphs[-1].text == "Before mark past mark"


# ---- the guarantees --------------------------------------------------------------------------------------------------------------------------------

def test_every_other_part_of_the_file_is_byte_for_byte_the_same_and_the_input_is_never_modified():
    data = _fixture()
    snapshot = bytes(data)
    edited = de.apply_edits(data, [{"op": "replace", "paragraph": 3, "old": "React", "new": "Vue"}])
    before, after = _members(data), _members(edited)
    assert list(before) == list(after)
    assert {name for name in before if before[name] != after[name]} == {"word/document.xml"}
    assert data == snapshot                              # the original bytes are untouched


def test_untouched_paragraphs_tables_and_settings_are_identical_in_the_copy():
    data = _fixture()
    edited = de.apply_edits(data, [{"op": "replace", "paragraph": 3, "old": "React", "new": "Vue"}])
    old_root, new_root = (etree.fromstring(_members(d)["word/document.xml"]) for d in (data, edited))
    old_paragraphs, new_paragraphs = list(old_root.iter(_q("p"))), list(new_root.iter(_q("p")))
    assert len(old_paragraphs) == len(new_paragraphs)
    for number, (old, new) in enumerate(zip(old_paragraphs, new_paragraphs), start=1):
        if number != 3:
            assert etree.tostring(old, method="c14n") == etree.tostring(new, method="c14n"), number
    assert de._skeleton(old_root) == de._skeleton(new_root)
    header = [n for n in _members(data) if "header" in n][0]
    assert _members(data)[header] == _members(edited)[header]          # the header was never in play


def test_the_safety_check_refuses_a_copy_where_something_else_changed(monkeypatch):
    real = de._replace_in_paragraph

    def sloppy(paragraph, old, new):
        real(paragraph, old, new)
        sibling = paragraph.getnext()
        if sibling is not None and sibling.find(_q("r")) is not None:
            sibling.find(_q("r")).find(_q("t")).text = "SNUCK IN"

    monkeypatch.setattr(de, "_replace_in_paragraph", sloppy)
    with pytest.raises(de.DocxEditError, match="did not pass its safety check"):
        de.apply_edits(_fixture(), [{"op": "replace", "paragraph": 2, "old": "Software", "new": "Hardware"}])


def test_the_safety_check_refuses_a_copy_where_a_paragraph_lost_its_formatting(monkeypatch):
    real = de._replace_in_paragraph

    def sloppy(paragraph, old, new):
        real(paragraph, old, new)
        properties = paragraph.find(_q("pPr"))
        if properties is not None:
            paragraph.remove(properties)

    monkeypatch.setattr(de, "_replace_in_paragraph", sloppy)
    with pytest.raises(de.DocxEditError, match="did not pass its safety check"):
        de.apply_edits(_fixture(), [{"op": "replace", "paragraph": 8, "old": "Sonic", "new": "Max"}])


def test_the_edited_copy_opens_as_a_normal_document_and_a_second_edit_can_be_made_to_it():
    first = de.apply_edits(_fixture(), [{"op": "replace", "paragraph": 3, "old": "React", "new": "Vue"}])
    second = de.apply_edits(first, [{"op": "replace", "paragraph": 3, "old": "Vue", "new": "Svelte"}])
    assert "Skills: Python, FastAPI, Svelte." in _texts(second)


# ---- planning errors the model can act on -------------------------------------------------------------------------------------------------------

@pytest.mark.parametrize("ops, fragment", [
    ([], "non-empty list"),
    ([{"op": "explode", "paragraph": 1}], "op must be"),
    ([{"op": "replace", "paragraph": 99, "old": "a", "new": "b"}], "no paragraph 99"),
    ([{"op": "replace", "paragraph": 3, "old": "nope", "new": "b"}], "not found in paragraph 3"),
    ([{"op": "replace", "paragraph": 3, "old": "a", "new": "b"}], "appears 2 times"),           # "Skills: Python, FastAPI, React." has two lower-case a's
    ([{"op": "replace", "paragraph": 3, "old": "Python", "new": "Python"}], "identical"),
    ([{"op": "replace", "paragraph": 3, "old": "", "new": "b"}], "is empty"),
    ([{"op": "replace", "paragraph": 3, "old": "Python", "new": "x" * 6000}], "too long"),
    ([{"op": "replace", "paragraph": 3, "old": "Python", "new": "Go"}, {"op": "delete", "paragraph": 3}], "already changed"),
    ([{"op": "insert_after", "paragraph": 3, "text": ""}], "is empty"),
    ([{"op": "delete", "paragraph": 3.5}], "no paragraph 3.5"),
])
def test_bad_edits_are_refused_with_a_message_that_says_what_to_fix(ops, fragment):
    with pytest.raises(de.DocxEditError, match=fragment):
        de.apply_edits(_fixture(), ops)


def test_at_most_forty_edits_at_once():
    ops = [{"op": "insert_after", "paragraph": 1, "text": f"line {i}"} for i in range(de.MAX_OPS + 1)]
    with pytest.raises(de.DocxEditError, match="too many word edits"):
        de.apply_edits(_fixture(), ops)


def test_the_same_planning_rules_work_on_the_outline_that_sonic_reads():
    outline = _outline(_fixture())
    plan = de.plan_edits(outline, [{"op": "replace", "paragraph": 3, "old": "React", "new": "Vue"}, {"op": "delete", "paragraph": 9}])
    assert de.edited_texts(outline, plan)[2] == "Skills: Python, FastAPI, Vue."
    assert len(de.edited_texts(outline, plan)) == len(de.original_texts(outline)) - 1
    with pytest.raises(de.DocxEditError, match="can't be deleted"):
        de.plan_edits(outline, [{"op": "delete", "paragraph": 4}])


# ---- through the folder snapshot and the edit proposal --------------------------------------------------------------------------------------------

import base64

from backend.services import local_edits as le
from backend.services import local_workspace as lw


def _b64(data):
    return base64.b64encode(data).decode("ascii")


def test_a_document_sent_to_the_snapshot_is_kept_only_as_its_outline_never_as_the_file(tmp_path, monkeypatch):
    monkeypatch.setattr(lw, "WORKSPACE_BASE_DIR", tmp_path)
    data = _fixture()
    result = lw.apply_sync("jack@example.com", "Docs", [{"path": "Resume.docx", "content_base64": _b64(data)}, {"path": "notes.txt", "content": "hi"}], [], True)
    assert result["accepted"] == 2 and result["rejected"] == []
    stored = next(tmp_path.rglob("Resume.docx")).read_bytes()
    assert stored.startswith(b"WORD DOCUMENT") and b"PK" not in stored[:4] and data not in stored
    assert de.parse_outline(stored.decode("utf-8"))[2]["text"] == "Software engineer at Principal since 2023."


@pytest.mark.parametrize("entry, reason", [
    ({"path": "Resume.docx", "content_base64": "!!!not base64!!!"}, "could not be read"),
    ({"path": "Resume.docx", "content_base64": _b64(b"plain text, not a zip")}, "valid .docx"),
    ({"path": "notes.txt", "content_base64": _b64(b"x")}, "only .docx"),
])
def test_a_document_that_cannot_be_used_is_rejected_with_a_reason_and_the_rest_of_the_batch_still_goes_through(tmp_path, monkeypatch, entry, reason):
    monkeypatch.setattr(lw, "WORKSPACE_BASE_DIR", tmp_path)
    result = lw.apply_sync("jack@example.com", "Docs", [entry, {"path": "ok.txt", "content": "fine"}], [], True)
    assert result["accepted"] == 1 and reason in result["rejected"][0]["reason"]


def _proposal(edits, outline_of=None):
    files = {"Resume.docx": de.outline_text(outline_of or _fixture())}
    return le.validate_edit_proposal(edits, lambda path: files.get(path))


def test_an_edit_to_an_existing_document_becomes_a_card_showing_the_text_before_and_after_and_saves_as_a_new_name():
    result = _proposal([{"path": "Resume.docx", "word_edits": [{"op": "replace", "paragraph": 3, "old": "React", "new": "Vue"}]}])
    assert result["ok"]
    assert result["edits"] == [{"type": "edit_docx", "path": "Resume.docx", "save_as": "Resume (edited).docx",
                                "word_edits": [{"op": "replace", "paragraph": 3, "old": "React", "new": "Vue"}]}]
    card = result["files"][0]
    assert card["path"] == "Resume (edited).docx" and card["kind"] == "create" and card["format"] == "docx_edit" and card["source"] == "Resume.docx"
    assert "-Skills: Python, FastAPI, React." in card["diff"] and "+Skills: Python, FastAPI, Vue." in card["diff"]
    assert "a/Resume.docx" in card["diff"] and card["additions"] == 1 and card["deletions"] == 1


def test_the_copy_name_can_be_chosen_but_never_the_original_never_an_existing_file_never_a_protected_path():
    ok = _proposal([{"path": "Resume.docx", "save_as": "drafts/Resume v2.docx", "word_edits": [{"op": "delete", "paragraph": 9}]}])
    assert ok["ok"] and ok["edits"][0]["save_as"] == "drafts/Resume v2.docx"
    for save_as, fragment in (("Resume.docx", "must differ"), ("resume.DOCX", "must differ"), ("Other.txt", "ending in .docx"),
                              ("../x.docx", "ending in .docx"), ("node_modules/x.docx", "not allowed")):
        bad = _proposal([{"path": "Resume.docx", "save_as": save_as, "word_edits": [{"op": "delete", "paragraph": 9}]}])
        assert not bad["ok"] and fragment in bad["errors"][0], save_as


def test_a_document_that_was_not_read_from_the_folder_cannot_be_edited_and_the_error_says_how_to_fix_it():
    result = le.validate_edit_proposal([{"path": "Missing.docx", "word_edits": [{"op": "delete", "paragraph": 1}]}], lambda path: None)
    assert not result["ok"] and "Read it first" in result["errors"][0]
    text = le.validate_edit_proposal([{"path": "notes.txt", "word_edits": [{"op": "delete", "paragraph": 1}]}], lambda path: "hi")
    assert not text["ok"] and "only apply to a .docx" in text["errors"][0]


def test_edits_that_cannot_be_applied_to_the_outline_are_refused_at_proposal_time_with_the_same_reasons():
    result = _proposal([{"path": "Resume.docx", "word_edits": [{"op": "replace", "paragraph": 3, "old": "Rust", "new": "Go"}]}])
    assert not result["ok"] and "not found in paragraph 3" in result["errors"][0]
    result = _proposal([{"path": "Resume.docx", "word_edits": [{"op": "delete", "paragraph": 6}]}])
    assert not result["ok"] and "only paragraph in its table cell" in result["errors"][0]


def test_a_second_edit_to_the_same_copy_name_in_one_proposal_is_refused_and_so_is_editing_a_document_this_proposal_creates():
    edit = {"path": "Resume.docx", "word_edits": [{"op": "delete", "paragraph": 9}]}
    twice = _proposal([edit, edit])
    assert not twice["ok"] and "already exists" in twice["errors"][0]
    doc = {"blocks": [{"type": "paragraph", "text": "x"}]}
    chained = _proposal([{"path": "New.docx", "document": doc}, {"path": "New.docx", "word_edits": [{"op": "delete", "paragraph": 1}]}])
    assert not chained["ok"] and "created by this same proposal" in chained["errors"][0]


def test_the_edit_endpoint_needs_sign_in_refuses_guests_and_returns_the_verified_copy_only():
    source = open("app.py", encoding="utf-8").read()
    start = source.index('@app.post("/api/local-edits/edit-docx")')
    endpoint = source[start:start + 1900]
    assert "Depends(get_current_user)" in endpoint and "is_guest_username" in endpoint and "status_code=403" in endpoint
    assert "apply_docx_edits" in endpoint and "except DocxEditError" in endpoint and "status_code=400" in endpoint and "DOCX_MIME" in endpoint
    assert "b64decode(payload.docx_base64, validate=True)" in endpoint
