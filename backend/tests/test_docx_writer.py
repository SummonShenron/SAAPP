"""Word documents Sonic writes into the connected folder (backend/utils/docx_writer.py, and the {path, document} edit in backend/services/local_edits.py)."""
import io
import zipfile

import pytest
from docx import Document

from backend.services import local_edits as le
from backend.utils import docx_writer as dw

RESUME = {
    "title": "Jack Harper",
    "blocks": [
        {"type": "heading", "text": "Experience", "level": 1},
        {"type": "paragraph", "text": "Built **Sonic**, an assistant that remembers."},
        {"type": "bullets", "items": ["Backend on FastAPI", "Frontend on React"]},
        {"type": "numbered", "items": ["First", "Second"]},
        {"type": "table", "rows": [["Role", "Years"], ["Engineer", "3"]], "header": True},
    ],
}


def _clean(spec=RESUME):
    clean, errors = dw.validate_document(spec)
    assert errors == [] and clean is not None
    return clean


# ---- checking the description -----------------------------------------------------------------------------------------------------------------

def test_a_good_document_is_accepted_and_normalised():
    clean = _clean()
    assert clean["title"] == "Jack Harper" and [b["type"] for b in clean["blocks"]] == ["heading", "paragraph", "bullets", "numbered", "table"]
    assert clean["blocks"][4]["header"] is True


@pytest.mark.parametrize("spec, fragment", [
    (None, "object"),
    ({}, "non-empty list of blocks"),
    ({"blocks": []}, "non-empty list of blocks"),
    ({"blocks": [{"type": "image", "text": "x"}]}, "type must be one of"),
    ({"blocks": [{"type": "paragraph", "text": 5}]}, "must be text"),
    ({"blocks": [{"type": "bullets", "items": []}]}, "non-empty list"),
    ({"blocks": [{"type": "table", "rows": [[]]}]}, "rows must be"),
    ({"blocks": [{"type": "paragraph", "text": "x" * (dw.MAX_TEXT + 1)}]}, "too long"),
    ({"blocks": [{"type": "paragraph", "text": "x"}] * (dw.MAX_BLOCKS + 1)}, "too many blocks"),
    ({"blocks": [{"type": "table", "rows": [["a"] * (dw.MAX_TABLE_COLS + 1)]}]}, "at most"),
])
def test_a_bad_description_comes_back_as_readable_problems_never_an_exception(spec, fragment):
    clean, errors = dw.validate_document(spec)
    assert clean is None and any(fragment in error for error in errors), errors


def test_the_whole_document_has_a_size_limit_too():
    blocks = [{"type": "paragraph", "text": "x" * 4000} for _ in range(10)]
    clean, errors = dw.validate_document({"blocks": blocks})
    assert clean is None and any("document is too long" in error for error in errors)


def test_a_heading_level_outside_one_to_three_becomes_one_and_a_ragged_table_is_padded():
    clean = _clean({"blocks": [{"type": "heading", "text": "A **b**", "level": 9}, {"type": "table", "rows": [["a", "b"], ["c"]]}]})
    assert clean["blocks"][0] == {"type": "heading", "text": "A b", "level": 1}
    assert clean["blocks"][1]["rows"] == [["a", "b"], ["c", ""]]


def test_control_characters_that_xml_cannot_carry_are_removed():
    clean = _clean({"blocks": [{"type": "paragraph", "text": "ok\x00\x08\x0bfine"}]})
    assert clean["blocks"][0]["text"] == "okfine"
    assert dw.render_docx(clean)       # python-docx would have raised on them


# ---- the preview the person reads before anything is written ----------------------------------------------------------------------------------

def test_the_preview_reads_like_the_document():
    preview = dw.preview_text(_clean())
    assert preview.startswith("# Jack Harper\n")
    assert "## Experience" in preview                       # headings sit under the title
    assert "- Backend on FastAPI" in preview and "1. First" in preview and "2. Second" in preview
    assert "| Role | Years |" in preview and "|---|---|" in preview and "| Engineer | 3 |" in preview


# ---- the real file ----------------------------------------------------------------------------------------------------------------------------

def test_the_rendered_file_is_a_real_word_document_with_the_same_content():
    data = dw.render_docx(_clean())
    assert zipfile.is_zipfile(io.BytesIO(data)) and data[:2] == b"PK"
    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    assert "Jack Harper" in texts and "Experience" in texts and "Backend on FastAPI" in texts and "Second" in texts
    paragraph = next(p for p in doc.paragraphs if p.text.startswith("Built "))
    assert paragraph.text == "Built Sonic, an assistant that remembers."
    assert [run.bold for run in paragraph.runs if run.text == "Sonic"] == [True]      # **bold** became bold, and the markers are gone
    table = doc.tables[0]
    assert [[cell.text for cell in row.cells] for row in table.rows] == [["Role", "Years"], ["Engineer", "3"]]
    assert all(run.bold for run in table.rows[0].cells[0].paragraphs[0].runs)         # the header row is bold
    styles = {p.style.name for p in doc.paragraphs}
    assert {"Title", "Heading 1", "List Bullet", "List Number"} <= styles


def test_nothing_in_the_file_can_run_code():
    names = zipfile.ZipFile(io.BytesIO(dw.render_docx(_clean()))).namelist()
    assert not any("vba" in name.lower() or name.endswith(".bin") for name in names)


# ---- as an edit proposal ----------------------------------------------------------------------------------------------------------------------

def _validate(edits, files=None):
    files = files or {}
    return le.validate_edit_proposal(edits, lambda path: files.get(path))


def test_a_new_word_document_becomes_a_card_with_a_readable_preview_and_a_normalised_edit():
    result = _validate([{"path": "Resume.docx", "document": RESUME}])
    assert result["ok"]
    assert result["edits"][0]["type"] == "create_docx" and result["edits"][0]["path"] == "Resume.docx"
    card = result["files"][0]
    assert card["kind"] == "create" and card["format"] == "docx" and card["additions"] > 5 and card["deletions"] == 0
    assert "+# Jack Harper" in card["diff"] and "+- Backend on FastAPI" in card["diff"]


def test_only_dotdocx_paths_new_names_and_one_edit_per_document_are_allowed():
    assert "ending in .docx" in _validate([{"path": "Resume.txt", "document": RESUME}])["errors"][0]
    assert "already exists" in _validate([{"path": "Resume.docx", "document": RESUME}], {"Resume.docx": "old"})["errors"][0]
    twice = _validate([{"path": "A.docx", "document": RESUME}, {"path": "A.docx", "old_string": "x", "new_string": "y"}])
    assert not twice["ok"] and "created by this same proposal" in twice["errors"][0]


def test_a_secret_path_or_a_bad_document_is_refused_before_anything_is_proposed():
    assert not _validate([{"path": "../Resume.docx", "document": RESUME}])["ok"]
    bad = _validate([{"path": "Resume.docx", "document": {"blocks": [{"type": "image"}]}}])
    assert not bad["ok"] and "edit 1 (Resume.docx): block 1" in bad["errors"][0]


def test_a_word_document_can_ride_with_ordinary_text_edits_in_one_proposal():
    result = _validate([{"path": "notes.txt", "content": "hello"}, {"path": "Letter.docx", "document": {"blocks": [{"type": "paragraph", "text": "Hi"}]}}])
    assert result["ok"] and [f["path"] for f in result["files"]] == ["notes.txt", "Letter.docx"]
    assert [f["format"] for f in result["files"]] == ["text", "docx"]


def test_the_render_endpoint_needs_sign_in_refuses_guests_and_bad_documents():
    source = open("app.py", encoding="utf-8").read()
    start = source.index('@app.post("/api/local-edits/render-docx")')
    endpoint = source[start:start + 1500]
    assert "Depends(get_current_user)" in endpoint and "is_guest_username" in endpoint and "status_code=403" in endpoint
    assert "validate_document" in endpoint and "status_code=400" in endpoint and "asyncio.to_thread(render_docx" in endpoint
    assert "DOCX_MIME" in endpoint
