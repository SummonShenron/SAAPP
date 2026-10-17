"""Word documents Sonic writes for the user's connected folder (see backend/services/local_edits.py, propose_local_edits).

Sonic never produces file bytes. It describes a document in a small structured form (headings, paragraphs, bullet and numbered lists, tables), this module checks
that form, shows it as a readable preview for the diff card, and renders the real .docx only when the user presses Apply (the browser asks for it, then
writes the file itself). A fixed set of block types, with size limits, keeps what gets written predictable: no macros, no images, no links, no raw XML.

    {"title": "Optional title", "blocks": [
        {"type": "heading", "text": "Experience", "level": 1},
        {"type": "paragraph", "text": "Plain text, with **bold** where needed."},
        {"type": "bullets", "items": ["one", "two"]},
        {"type": "numbered", "items": ["first", "second"]},
        {"type": "table", "rows": [["Name", "Role"], ["Ada", "Engineer"]], "header": true}]}

Pure functions apart from the python-docx rendering; never raise on bad input (problems come back as readable strings for the model to fix).
"""
import io
import re
from typing import Any, Dict, List, Optional, Tuple

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_BLOCKS = 200
MAX_CHARS = 30_000
MAX_TEXT = 5_000
MAX_ITEMS = 100
MAX_TABLE_ROWS = 50
MAX_TABLE_COLS = 10
MAX_TITLE = 120
BLOCK_TYPES = ("heading", "paragraph", "bullets", "numbered", "table")

# XML 1.0 cannot carry these characters and python-docx refuses them.
_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")
_BOLD = re.compile(r"\*\*(.+?)\*\*", re.DOTALL)


def _clean(value: Any, limit: int, what: str, errors: List[str]) -> Optional[str]:
    if not isinstance(value, str):
        errors.append(f"{what} must be text")
        return None
    text = _CONTROL_CHARS.sub("", value.replace("\r\n", "\n")).strip("\n")
    if len(text) > limit:
        errors.append(f"{what} is too long ({len(text)} characters; the limit is {limit})")
        return None
    return text


def _items(raw: Any, label: str, errors: List[str]) -> Optional[List[str]]:
    if not isinstance(raw, list) or not raw:
        errors.append(f"{label}: items must be a non-empty list of text")
        return None
    if len(raw) > MAX_ITEMS:
        errors.append(f"{label}: too many items ({len(raw)}; the limit is {MAX_ITEMS})")
        return None
    cleaned = [_clean(item, MAX_TEXT, f"{label}: an item", errors) for item in raw]
    return None if any(item is None for item in cleaned) else [item for item in cleaned if item is not None]


def validate_document(spec: Any) -> Tuple[Optional[Dict[str, Any]], List[str]]:
    """(clean document, []) or (None, [problems]). The clean form is what the preview and the renderer both use."""
    if not isinstance(spec, dict):
        return None, ["document must be an object with a list of blocks"]
    errors: List[str] = []
    title = ""
    if spec.get("title") not in (None, ""):
        title = _clean(spec.get("title"), MAX_TITLE, "title", errors) or ""
    blocks = spec.get("blocks")
    if not isinstance(blocks, list) or not blocks:
        return None, errors + ["document needs a non-empty list of blocks"]
    if len(blocks) > MAX_BLOCKS:
        return None, errors + [f"too many blocks ({len(blocks)}; the limit is {MAX_BLOCKS})"]

    clean_blocks: List[Dict[str, Any]] = []
    for index, block in enumerate(blocks, start=1):
        label = f"block {index}"
        if not isinstance(block, dict) or block.get("type") not in BLOCK_TYPES:
            errors.append(f"{label}: type must be one of {', '.join(BLOCK_TYPES)}")
            continue
        kind = block["type"]
        if kind in ("heading", "paragraph"):
            text = _clean(block.get("text"), MAX_TEXT, f"{label} text", errors)
            if text is None:
                continue
            if kind == "heading":
                level = block.get("level", 1)
                level = level if isinstance(level, int) and not isinstance(level, bool) and 1 <= level <= 3 else 1
                clean_blocks.append({"type": "heading", "text": text.replace("**", ""), "level": level})
            else:
                clean_blocks.append({"type": "paragraph", "text": text})
        elif kind in ("bullets", "numbered"):
            items = _items(block.get("items"), label, errors)
            if items is not None:
                clean_blocks.append({"type": kind, "items": items})
        else:  # table
            rows = block.get("rows")
            if not isinstance(rows, list) or not rows or not all(isinstance(row, list) and row for row in rows):
                errors.append(f"{label}: rows must be a non-empty list of non-empty lists of text")
                continue
            if len(rows) > MAX_TABLE_ROWS or max(len(row) for row in rows) > MAX_TABLE_COLS:
                errors.append(f"{label}: a table may have at most {MAX_TABLE_ROWS} rows and {MAX_TABLE_COLS} columns")
                continue
            width = max(len(row) for row in rows)
            cells = [[_clean(cell, MAX_TEXT, f"{label}: a cell", errors) for cell in row] for row in rows]
            if any(cell is None for row in cells for cell in row):
                continue
            padded = [[cell or "" for cell in row] + [""] * (width - len(row)) for row in cells]
            clean_blocks.append({"type": "table", "rows": padded, "header": block.get("header", True) is not False})

    total = len(title) + sum(len(" ".join(_block_texts(block))) for block in clean_blocks)
    if not errors and total > MAX_CHARS:
        errors.append(f"the document is too long ({total} characters; the limit is {MAX_CHARS})")
    if errors:
        return None, errors[:12]
    return {"title": title, "blocks": clean_blocks}, []


def _block_texts(block: Dict[str, Any]) -> List[str]:
    if block["type"] in ("heading", "paragraph"):
        return [block["text"]]
    if block["type"] in ("bullets", "numbered"):
        return list(block["items"])
    return [cell for row in block["rows"] for cell in row]


def preview_text(document: Dict[str, Any]) -> str:
    """The document as readable plain text (a light markdown), which is what the diff card shows before anything is written."""
    lines: List[str] = []
    if document.get("title"):
        lines += [f"# {document['title']}", ""]
    for block in document["blocks"]:
        kind = block["type"]
        if kind == "heading":
            lines += ["#" * (block["level"] + (1 if document.get("title") else 0)) + " " + block["text"], ""]
        elif kind == "paragraph":
            lines += [block["text"], ""]
        elif kind == "bullets":
            lines += [f"- {item}" for item in block["items"]] + [""]
        elif kind == "numbered":
            lines += [f"{number}. {item}" for number, item in enumerate(block["items"], start=1)] + [""]
        else:
            for position, row in enumerate(block["rows"]):
                lines.append("| " + " | ".join(cell.replace("\n", " ") for cell in row) + " |")
                if position == 0 and block["header"]:
                    lines.append("|" + "---|" * len(row))
            lines.append("")
    return "\n".join(lines).rstrip("\n") + "\n"


def _add_runs(paragraph: Any, text: str) -> None:
    """Adds `text` to the paragraph, bold where it was written **like this**."""
    position = 0
    for match in _BOLD.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        paragraph.add_run(match.group(1)).bold = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def render_docx(document: Dict[str, Any]) -> bytes:
    """The .docx bytes for a document that validate_document accepted."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    if document.get("title"):
        doc.add_heading(document["title"], level=0)
    for block in document["blocks"]:
        kind = block["type"]
        if kind == "heading":
            doc.add_heading(block["text"], level=block["level"])
        elif kind == "paragraph":
            _add_runs(doc.add_paragraph(), block["text"])
        elif kind in ("bullets", "numbered"):
            style = "List Bullet" if kind == "bullets" else "List Number"
            for item in block["items"]:
                _add_runs(doc.add_paragraph(style=style), item)
        else:
            rows = block["rows"]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                for col_index, cell_text in enumerate(row):
                    paragraph = table.cell(row_index, col_index).paragraphs[0]
                    _add_runs(paragraph, cell_text)
                    if row_index == 0 and block["header"]:
                        for run in paragraph.runs:
                            run.bold = True
            doc.add_paragraph()
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
