"""Plain text, Markdown and Word uploads: read correctly, refused in words when unusable, and a failure leaves nothing behind (local_function_app/text_files.py, function_app.py)."""
import io
import os
import zipfile

import pytest

from local_function_app import function_app as fa
from local_function_app import text_files as tf
from local_function_app.tests.test_function_app_json import FakeFile, FakeFS, pipeline  # noqa: F401  (the fixture is used by name)

SENTENCE = "The retry delay is two seconds and the job gives up after three tries. "


def make_docx(build):
    from docx import Document as DocxDocument

    document = DocxDocument()
    build(document)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


class Recorder:
    def __init__(self):
        self.edge_deletes = []

    def __call__(self, db, affiliate, source):
        self.edge_deletes.append((affiliate, source))


@pytest.fixture
def edges(monkeypatch):
    recorder = Recorder()
    monkeypatch.setattr(fa, "delete_graph_edges", recorder)
    return recorder


# ---- reading the files ----------------------------------------------------------------------------------------------------------------------

def test_the_file_names_are_recognised():
    assert all(fa._is_text_filename(n) for n in ("a.txt", "SERVER.LOG", "x.Txt"))
    assert all(fa._is_markdown_filename(n) for n in ("README.md", "notes.MARKDOWN"))
    assert all(fa._is_docx_filename(n) for n in ("Plan.docx", "PLAN.DOCX"))
    assert not any(fa._is_text_filename(n) or fa._is_markdown_filename(n) or fa._is_docx_filename(n) for n in ("a.pdf", "week.json", "photo.png", "txt", "doc.doc"))


def test_text_is_read_as_the_encoding_it_is():
    assert tf.decode_text("héllo wörld".encode("utf-8")) == "héllo wörld"
    assert tf.decode_text(b"\xef\xbb\xbfplain with a BOM") == "plain with a BOM"
    assert tf.decode_text("café au lait".encode("utf-16")) == "café au lait"                                   # Notepad's "Unicode"
    assert tf.decode_text("café au lait".encode("cp1252")) == "café au lait"                                   # Windows ANSI is not valid UTF-8
    assert tf.decode_text(b"line one\r\nline two\rline three\x00") == "line one\nline two\nline three"


def test_binary_empty_and_oversized_files_are_refused_in_words(monkeypatch):
    for raw, fragment in ((bytes(range(256)) * 40, "binary"), (b"   \n\n  ", "no readable text"), (b"", "no readable text")):
        with pytest.raises(tf.TextIngestError, match=fragment):
            tf.build_text_document(raw, "a.txt", "personal_jack", fa.Document, "i1")
    monkeypatch.setattr(tf, "MAX_BYTES", 100)
    with pytest.raises(tf.TextIngestError, match="over the 0 MB limit|limit for text files"):
        tf.build_text_document(b"x" * 101, "a.log", "personal_jack", fa.Document, "i1")


def test_a_word_document_is_read_in_order_with_tables_where_they_sit_and_merged_cells_once():
    def build(doc):
        doc.add_paragraph("Intro paragraph")
        table = doc.add_table(rows=2, cols=3)
        for index, text in enumerate(("Name", "Role", "Team")):
            table.cell(0, index).text = text
        merged = table.cell(1, 0).merge(table.cell(1, 1))
        merged.text = "Dana (lead)"
        table.cell(1, 2).text = "Platform"
        doc.add_paragraph("After the table")

    text = tf.docx_text(make_docx(build))
    assert text == "Intro paragraph\n\nName | Role | Team\nDana (lead) | Platform\n\nAfter the table"


def test_a_word_file_that_is_not_one_empty_or_huge_is_refused_in_words(monkeypatch):
    with pytest.raises(tf.TextIngestError, match="could not be read as a Word document"):
        tf.docx_text(b"this is not a zip file")
    with pytest.raises(tf.TextIngestError, match="no readable text"):
        tf.build_docx_document(make_docx(lambda doc: None), "e.docx", "personal_jack", fa.Document, "i1")
    monkeypatch.setattr(tf, "MAX_DOCX_UNZIPPED_BYTES", 10)
    with pytest.raises(tf.TextIngestError, match="expands to far more"):
        tf.docx_text(make_docx(lambda doc: doc.add_paragraph("some text that is longer than ten bytes")))
    monkeypatch.setattr(tf, "MAX_DOCX_BYTES", 10)
    with pytest.raises(tf.TextIngestError, match="limit for Word files"):
        tf.docx_text(make_docx(lambda doc: doc.add_paragraph("x")))


# ---- the pipeline: success ------------------------------------------------------------------------------------------------------------------

@pytest.mark.parametrize("filename,data,doc_type,content_type", [
    ("notes.txt", (SENTENCE * 30).encode(), "text", "text/plain"),
    ("server.log", (SENTENCE * 30).encode(), "text", "text/plain"),
    ("README.md", ("# Title\n\n" + SENTENCE * 30).encode(), "markdown", "text/markdown"),
])
def test_text_and_markdown_uploads_are_chunked_embedded_graphed_and_marked_done(pipeline, filename, data, doc_type, content_type):
    pipeline.fs = FakeFS([FakeFile("n1", filename, data)])
    assert fa.run_ingestion_pipeline() is True
    chunks = pipeline.stored[0]
    assert len(chunks) >= 2 and all(c.metadata["source"] == filename and c.metadata["doc_type"] == doc_type and c.metadata["affiliate"] == "personal_jack" for c in chunks)
    assert len({c.metadata["ingest_id"] for c in chunks}) == 1 and pipeline.graph == [("personal_jack", filename, len(chunks))]
    flt, update = pipeline.db["fs.files"].updates[0]
    assert flt == {"_id": "n1"} and update["$set"] == {"metadata.status": "pages", "metadata.content_type": content_type} and pipeline.fs.deleted == []


def test_a_word_upload_is_indexed_with_the_word_type(pipeline):
    data = make_docx(lambda doc: [doc.add_paragraph(SENTENCE * 5) for _ in range(6)])
    pipeline.fs = FakeFS([FakeFile("w1", "Plan.docx", data)])
    assert fa.run_ingestion_pipeline() is True
    assert all(c.metadata["doc_type"] == "docx" for c in pipeline.stored[0])
    assert pipeline.db["fs.files"].updates[0][1]["$set"]["metadata.content_type"].endswith("wordprocessingml.document")


def test_re_uploading_replaces_the_earlier_version_only_after_the_new_one_is_stored(pipeline):
    earlier = FakeFile("old1", "notes.txt", b"old", status="pages")
    other_library = FakeFile("other1", "notes.txt", b"old", affiliate="Affiliate_C", status="pages")
    pipeline.fs = FakeFS([FakeFile("new1", "notes.txt", (SENTENCE * 5).encode())], pages=[earlier, other_library])
    fa.run_ingestion_pipeline()
    new_id = pipeline.stored[0][0].metadata["ingest_id"]
    assert pipeline.db["documents"].deleted_queries == [{"source": "notes.txt", "affiliate": "personal_jack", "ingest_id": {"$ne": new_id}}]
    assert pipeline.fs.deleted == ["old1"]


def test_the_pdf_and_image_paths_never_see_these_files(pipeline, monkeypatch):
    monkeypatch.setattr(fa, "PyPDFLoader", lambda *a, **k: (_ for _ in ()).throw(AssertionError("must not reach the PDF loader")))
    pipeline.fs = FakeFS([FakeFile("a", "a.txt", (SENTENCE * 5).encode()), FakeFile("b", "b.md", (SENTENCE * 5).encode()),
                          FakeFile("c", "c.docx", make_docx(lambda doc: doc.add_paragraph(SENTENCE * 5)))])
    assert fa.run_ingestion_pipeline() is True and len(pipeline.stored) == 3


# ---- the pipeline: a failure leaves nothing behind -----------------------------------------------------------------------------------------

def test_an_unusable_file_is_wiped_completely_without_stopping_the_other_uploads_or_touching_an_earlier_version(pipeline, edges):
    earlier = FakeFile("old1", "empty.txt", b"old", status="pages")
    bad = FakeFile("bad1", "empty.txt", b"   ")
    good = FakeFile("good1", "good.txt", (SENTENCE * 5).encode())
    pipeline.fs = FakeFS([bad, good], pages=[earlier])
    assert fa.run_ingestion_pipeline() is True
    assert "bad1" in pipeline.fs.deleted and "good1" not in pipeline.fs.deleted and "old1" not in pipeline.fs.deleted
    assert len(pipeline.stored) == 1 and pipeline.stored[0][0].metadata["source"] == "good.txt"
    assert edges.edge_deletes == [], "nothing was written to the graph and the earlier version is still there, so its edges stay"
    chunk_queries = [q for q in pipeline.db["documents"].deleted_queries if q.get("source") == "empty.txt"]
    assert len(chunk_queries) == 1 and isinstance(chunk_queries[0]["ingest_id"], str)


def test_a_real_fault_before_anything_is_stored_wipes_this_upload_raises_and_leaves_the_earlier_version_whole(pipeline, edges):
    earlier = FakeFile("old1", "notes.txt", b"old", status="pages")
    pipeline.fs = FakeFS([FakeFile("new1", "notes.txt", (SENTENCE * 5).encode())], pages=[earlier])
    pipeline.fail_embedding = True
    with pytest.raises(RuntimeError):
        fa.run_ingestion_pipeline()
    assert pipeline.fs.deleted == ["new1"] and edges.edge_deletes == [] and pipeline.graph == []
    queries = pipeline.db["documents"].deleted_queries
    assert len(queries) == 1 and isinstance(queries[0]["ingest_id"], str)                                    # only this upload's own chunks, never the earlier version's


def test_a_fault_after_the_graph_step_removes_the_chunks_the_edges_and_the_stored_copy(pipeline, edges):
    pipeline.fs = FakeFS([FakeFile("new1", "notes.txt", (SENTENCE * 5).encode())])

    def boom(flt, update):
        raise RuntimeError("could not mark the file as done")

    pipeline.db["fs.files"].update_one = boom
    with pytest.raises(RuntimeError):
        fa.run_ingestion_pipeline()
    assert pipeline.graph and edges.edge_deletes == [("personal_jack", "notes.txt")] and pipeline.fs.deleted == ["new1"]
    last = pipeline.db["documents"].deleted_queries[-1]                      # the first query is the replace-earlier step; the last is this upload's own rollback
    assert last["source"] == "notes.txt" and isinstance(last["ingest_id"], str)


def test_a_fault_while_replacing_the_earlier_version_still_removes_the_edges_that_version_leaves_behind(pipeline, edges, monkeypatch):
    pipeline.fs = FakeFS([FakeFile("new1", "notes.txt", (SENTENCE * 5).encode())])
    monkeypatch.setattr(fa, "_replace_earlier_version", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("replace failed halfway")))
    with pytest.raises(RuntimeError):
        fa.run_ingestion_pipeline()
    assert edges.edge_deletes == [("personal_jack", "notes.txt")] and pipeline.fs.deleted == ["new1"] and pipeline.graph == []


def test_a_corrupt_word_file_and_a_file_with_too_many_chunks_are_wiped_and_refused(pipeline, edges, monkeypatch):
    monkeypatch.setattr(tf, "MAX_CHUNKS", 2)
    pipeline.fs = FakeFS([FakeFile("c1", "broken.docx", b"not a zip"), FakeFile("c2", "huge.log", (SENTENCE * 200).encode()), FakeFile("ok", "small.txt", b"short note, one chunk only")])
    assert fa.run_ingestion_pipeline() is True
    assert pipeline.fs.deleted == ["c1", "c2"] and len(pipeline.stored) == 1 and pipeline.stored[0][0].metadata["source"] == "small.txt" and edges.edge_deletes == []


# ---- the citation link ---------------------------------------------------------------------------------------------------------------------

def test_each_kind_of_file_is_served_as_what_it_is():
    from backend.utils import citation_media as cm

    assert cm.media_for("week.JSON") == ("application/json", "inline") and cm.media_for("a.jsonl") == ("application/json", "inline")
    for name in ("notes.txt", "server.LOG", "README.md", "notes.markdown"):
        assert cm.media_for(name) == ("text/plain; charset=utf-8", "inline")
    assert cm.media_for("data.CSV") == ("text/plain; charset=utf-8", "inline") and cm.media_for("t.tsv") == ("text/plain; charset=utf-8", "inline")
    assert cm.media_for("book.xlsx") == (cm.XLSX_TYPE, "attachment") and cm.media_for("book.xlsm")[1] == "attachment"
    assert cm.media_for("Plan.docx") == (cm.DOCX_TYPE, "attachment") and cm.media_for("report.pdf") == ("application/pdf", "inline") and cm.media_for("") == ("application/pdf", "inline")
    source = open(os.path.join(os.path.dirname(__file__), "..", "..", "app.py"), encoding="utf-8").read()
    assert "media_type, disposition = citation_media.media_for(clean_basename)" in source and 'f"{disposition}; filename*=' in source


def test_the_run_never_prints_the_database_user_or_password():
    assert fa._redacted_mongo_host("mongodb+srv://someuser:S3cret%21pass@cluster0.abc.mongodb.net/?appName=App") == "cluster0.abc.mongodb.net"
    assert fa._redacted_mongo_host("mongodb://localhost:27017") == "localhost:27017" and fa._redacted_mongo_host(None) == "(not set)"
    source = open(os.path.join(os.path.dirname(__file__), "..", "function_app.py"), encoding="utf-8").read()
    assert "MONGO_URI is currently" not in source


# ---- spreadsheets (.csv, .tsv, .xlsx, .xlsm) ---------------------------------------------------------------------------------------------------

from local_function_app import spreadsheet_files as sf  # noqa: E402


def make_xlsx(sheets, shared=("Name", "Role", "Dana", "Lead")):
    """A small real workbook: `sheets` is {name: [[cell, ...], ...]} where a cell is a number, a string stored inline, or ("s", index) into the shared strings."""
    def cell_xml(row, col, value):
        ref = f"{chr(ord('A') + col)}{row}"
        if isinstance(value, tuple):
            return f'<c r="{ref}" t="s"><v>{value[1]}</v></c>'
        if isinstance(value, str):
            return f'<c r="{ref}" t="inlineStr"><is><t>{value}</t></is></c>'
        return f'<c r="{ref}"><v>{value}</v></c>'

    ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("xl/workbook.xml", f'<workbook {ns} xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>'
                   + "".join(f'<sheet name="{name}" sheetId="{n}" r:id="rId{n}"/>' for n, name in enumerate(sheets, 1)) + "</sheets></workbook>")
        z.writestr("xl/_rels/workbook.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   + "".join(f'<Relationship Id="rId{n}" Target="worksheets/sheet{n}.xml"/>' for n in range(1, len(sheets) + 1)) + "</Relationships>")
        z.writestr("xl/sharedStrings.xml", f"<sst {ns}>" + "".join(f"<si><t>{t}</t></si>" for t in shared) + "</sst>")
        for n, rows in enumerate(sheets.values(), 1):
            body = "".join(f'<row r="{r}">' + "".join(cell_xml(r, c, v) for c, v in enumerate(row)) + "</row>" for r, row in enumerate(rows, 1))
            z.writestr(f"xl/worksheets/sheet{n}.xml", f"<worksheet {ns}><sheetData>{body}</sheetData></worksheet>")
    return out.getvalue()


def test_csv_and_tsv_are_read_as_tables_in_any_encoding():
    assert sf.spreadsheet_text(b'Name,Role\r\n"Lopez, Dana",Lead\r\n,\r\n', "team.csv") == "Name | Role\nName: Lopez, Dana | Role: Lead"
    assert sf.spreadsheet_text("a\tb\n1\t2\n".encode("utf-16"), "t.TSV") == "a | b\na: 1 | b: 2"
    assert sf.spreadsheet_text("caf\u00e9,x\n".encode("cp1252"), "c.csv") == "caf\u00e9 | x"


def test_a_workbook_is_read_sheet_by_sheet_with_shared_and_inline_strings_and_numbers():
    raw = make_xlsx({"Team": [[("s", 0), ("s", 1)], [("s", 2), ("s", 3)], ["", ""]], "Hours": [["Week", 1], ["Total", 42.5]]})
    assert sf.spreadsheet_text(raw, "team.xlsx") == "Sheet: Team\nName | Role\nName: Dana | Role: Lead\n\nSheet: Hours\nWeek | 1\nTotal | 42.5"


def test_an_unreadable_empty_or_oversized_spreadsheet_is_refused_in_words(monkeypatch):
    for raw, name, fragment in ((b"not a zip", "a.xlsx", "not a readable Excel workbook"), (b"", "e.csv", "no readable cells"), (b",,\n,,\n", "e.csv", "no readable cells"),
                                (make_xlsx({"S": [[]]}), "e.xlsx", "no readable cells"), (b"x", "a.pdf", "Not a spreadsheet")):
        with pytest.raises(tf.TextIngestError, match=fragment):
            sf.spreadsheet_text(raw, name)
    monkeypatch.setattr(sf, "MAX_WORKBOOK_UNZIPPED_BYTES", 10)
    with pytest.raises(tf.TextIngestError, match="size limit"):
        sf.spreadsheet_text(make_xlsx({"S": [["some text that is long enough"]]}), "a.xlsx")
    monkeypatch.setattr(tf, "MAX_BYTES", 5)
    with pytest.raises(tf.TextIngestError, match="too big"):
        sf.spreadsheet_text(b"a,b,c,d,e,f", "big.csv")


SHEET_NS = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'


def raw_xlsx(sheets, hidden=(), styles=None, date1904=False):
    """A workbook built from raw <row> XML: `sheets` is {name: sheetData xml}; `hidden` names the sheets marked hidden; `styles` is the xl/styles.xml body (cellXfs, numFmts)."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("xl/workbook.xml", f'<workbook {SHEET_NS} xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                   + ('<workbookPr date1904="1"/>' if date1904 else "") + "<sheets>"
                   + "".join(f'<sheet name="{n}" sheetId="{i}"{" state=\"hidden\"" if n in hidden else ""} r:id="rId{i}"/>' for i, n in enumerate(sheets, 1)) + "</sheets></workbook>")
        z.writestr("xl/_rels/workbook.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   + "".join(f'<Relationship Id="rId{i}" Target="worksheets/sheet{i}.xml"/>' for i in range(1, len(sheets) + 1)) + "</Relationships>")
        if styles:
            z.writestr("xl/styles.xml", f"<styleSheet {SHEET_NS}>{styles}</styleSheet>")
        for i, xml in enumerate(sheets.values(), 1):
            z.writestr(f"xl/worksheets/sheet{i}.xml", f"<worksheet {SHEET_NS}><sheetData>{xml}</sheetData></worksheet>")
    return out.getvalue()


def text_cell(ref, text):
    return f'<c r="{ref}" t="inlineStr"><is><t>{text}</t></is></c>' if ref else f'<c t="inlineStr"><is><t>{text}</t></is></c>'


def test_every_row_after_the_header_says_which_column_each_value_is_in():
    big = "Name,Role,Hours\n" + "".join(f"Person{n},Lead,{n}\n" for n in range(1, 4))
    assert sf.spreadsheet_text(big.encode(), "t.csv") == "Name | Role | Hours\nName: Person1 | Role: Lead | Hours: 1\nName: Person2 | Role: Lead | Hours: 2\nName: Person3 | Role: Lead | Hours: 3"
    # a title above the header stays plain; an empty cell is left out and does not shift the labels; a blank header cell gives a bare value; a long header is cut
    table = b"Quarterly team sheet\n,,\nName,,Role\nDana,x,\nLee,,Lead\n"
    assert sf.spreadsheet_text(table, "t.csv").split("\n") == ["Quarterly team sheet", "Name |  | Role", "Name: Dana | x", "Name: Lee | Role: Lead"]
    assert sf.spreadsheet_text(("H" * 100 + ",b\nv,w\n").encode(), "t.csv").split("\n")[1] == "H" * 60 + ": v | b: w"
    # a first row with a number in it is data, not a header: nothing is labelled (so "Week | 1" never turns "Total | 42.5" into "Week: Total")
    assert sf.spreadsheet_text(b"2023,2024\nx,y\n", "t.csv") == "2023 | 2024\nx | y"
    assert sf.spreadsheet_text(b"only\none\ncolumn\n", "t.csv") == "only\none\ncolumn"
    # a cell with line breaks stays on one line
    assert sf.spreadsheet_text(b'Name,Notes\r\nDana,"line one\r\nline   two"\r\n', "t.csv").split("\n")[-1] == "Name: Dana | Notes: line one line two"


def test_a_cell_in_a_column_that_cannot_exist_is_refused_at_once_not_built():
    for ref in ("ZZZZZZZZ1", "XFE1", "AAAA1"):
        with pytest.raises(tf.TextIngestError, match="columns a spreadsheet can have"):
            sf.spreadsheet_text(raw_xlsx({"S": f'<row r="1"><c r="{ref}"><v>1</v></c></row>'}), "a.xlsx")
    assert sf.spreadsheet_text(raw_xlsx({"S": '<row r="1">' + text_cell("XFD1", "last") + "</row>"}), "a.xlsx").endswith("last")


def test_cells_without_positions_are_read_in_order_not_stacked_in_column_a():
    raw = raw_xlsx({"S": "<row>" + text_cell(None, "Name") + text_cell(None, "Role") + "</row><row>" + text_cell(None, "Dana") + text_cell(None, "Lead") + "</row>"})
    assert sf.spreadsheet_text(raw, "a.xlsx") == "Sheet: S\nName | Role\nName: Dana | Role: Lead"
    mixed = raw_xlsx({"S": "<row>" + text_cell("C1", "third") + text_cell(None, "fourth") + "</row>"})
    assert sf.spreadsheet_text(mixed, "a.xlsx") == "Sheet: S\n |  | third | fourth"


def test_hidden_sheets_are_not_read_and_a_workbook_of_only_hidden_sheets_is_refused():
    raw = raw_xlsx({"Public": "<row>" + text_cell("A1", "hello") + "</row>", "Secret": "<row>" + text_cell("A1", "salary 90000") + "</row>"}, hidden=("Secret",))
    text = sf.spreadsheet_text(raw, "a.xlsx")
    assert "hello" in text and "salary" not in text and "Secret" not in text
    with pytest.raises(tf.TextIngestError, match="no readable cells"):
        sf.spreadsheet_text(raw_xlsx({"Secret": "<row>" + text_cell("A1", "x") + "</row>"}, hidden=("Secret",)), "a.xlsx")


DATE_STYLES = ('<numFmts count="2"><numFmt numFmtId="164" formatCode="dd/mm/yyyy hh:mm"/><numFmt numFmtId="165" formatCode="0.00&quot;d&quot;E+00"/></numFmts>'
               '<cellXfs count="5"><xf numFmtId="0"/><xf numFmtId="14"/><xf numFmtId="164"/><xf numFmtId="165"/><xf numFmtId="21"/></cellXfs>')


def test_date_cells_are_written_as_dates_and_other_numbers_are_left_alone():
    sheet = ("<row>" + text_cell("A1", "When") + text_cell("B1", "At") + text_cell("C1", "Qty") + text_cell("D1", "Time") + "</row>"
             '<row><c r="A2" s="1"><v>45123</v></c><c r="B2" s="2"><v>45123.5625</v></c><c r="C2" s="3"><v>45123</v></c><c r="D2" s="4"><v>0.75</v></c></row>'
             '<row><c r="A3"><v>45123</v></c><c r="B3" s="1"><v>99999999</v></c></row>')
    text = sf.spreadsheet_text(raw_xlsx({"S": sheet}, styles=DATE_STYLES), "a.xlsx")
    assert "When: 2023-07-16 | At: 2023-07-16 13:30 | Qty: 45123 | Time: 18:00" in text            # a custom format with quoted text and E+ is not a date
    assert "When: 45123 | At: 99999999" in text                                                   # an unstyled number, and a serial past year 9999, stay numbers
    assert sf.spreadsheet_text(raw_xlsx({"S": '<row><c r="A1" s="1"><v>45123</v></c></row>'}, styles=DATE_STYLES, date1904=True), "a.xlsx").endswith("2027-07-17")
    assert sf.spreadsheet_text(raw_xlsx({"S": '<row><c r="A1" s="1"><v>1</v></c><c r="B1" s="1"><v>61</v></c></row>'}, styles=DATE_STYLES), "a.xlsx").endswith("1900-01-01 | 1900-03-01")
    assert sf._is_date_code("yyyy-mm-dd") and sf._is_date_code("[h]:mm:ss") and not sf._is_date_code("General") and not sf._is_date_code('0.00"days"') and not sf._is_date_code("[Red]0.0")


def test_a_formula_with_no_stored_value_is_left_empty_and_one_with_a_value_is_read():
    sheet = ('<row><c r="A1" t="inlineStr"><is><t>Total</t></is></c><c r="B1"><f>1+1</f></c><c r="C1" t="str"><f>"a"&amp;"b"</f><v>ab</v></c></row>')
    assert sf.spreadsheet_text(raw_xlsx({"S": sheet}), "a.xlsx") == "Sheet: S\nTotal |  | ab"


def test_the_error_class_is_the_one_the_pipeline_catches_however_the_folder_was_imported():
    assert sf.TextIngestError is tf.TextIngestError


@pytest.mark.parametrize("filename,data,content_type", [
    ("team.csv", b"Name,Role\n" + b"Dana,Lead\n" * 120, "text/csv"),
    ("team.tsv", b"Name\tRole\n" + b"Dana\tLead\n" * 120, "text/tab-separated-values"),
    ("team.xlsx", make_xlsx({"Team": [["Name", "Role"]] + [["Dana", "Lead"]] * 150}), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
])
def test_spreadsheets_go_through_the_shared_path_and_are_marked_with_their_own_type(pipeline, filename, data, content_type):
    pipeline.fs = FakeFS([FakeFile("s1", filename, data)])
    assert fa.run_ingestion_pipeline() is True
    assert all(c.metadata["doc_type"] == "spreadsheet" and c.metadata["source"] == filename for c in pipeline.stored[0])
    assert pipeline.db["fs.files"].updates[0][1]["$set"]["metadata.content_type"] == content_type and pipeline.fs.deleted == []


def test_a_bad_spreadsheet_is_wiped_like_any_unusable_file_and_the_others_carry_on(pipeline, edges):
    pipeline.fs = FakeFS([FakeFile("bad", "broken.xlsx", b"not a zip"), FakeFile("ok", "fine.csv", b"a,b\n" + b"1,2\n" * 200)])
    assert fa.run_ingestion_pipeline() is True
    assert pipeline.fs.deleted == ["bad"] and len(pipeline.stored) == 1 and edges.edge_deletes == []
    assert fa._is_spreadsheet_filename("A.XLSM") and not fa._is_spreadsheet_filename("a.xls") and not fa._is_spreadsheet_filename("a.pdf")
